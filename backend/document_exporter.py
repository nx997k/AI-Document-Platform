import re
from io import BytesIO

from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


def _escape_pdf(text):
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def _parse_blocks(content):
    blocks = []
    for block in content.split("\n\n"):
        block = block.strip()
        if block:
            blocks.append(block)
    return blocks


def _block_level(block):
    if block.startswith("### "):
        return 3, block[4:]
    if block.startswith("## "):
        return 2, block[3:]
    if block.startswith("# "):
        return 1, block[2:]
    return 0, block


def export_to_docx(content, title="Generated Document"):
    doc = Document()
    doc.add_heading(title, 0)

    for block in _parse_blocks(content):
        level, text = _block_level(block)

        if level == 1:
            doc.add_heading(text, level=1)
        elif level == 2:
            doc.add_heading(text, level=2)
        elif level == 3:
            doc.add_heading(text, level=3)
        elif re.match(r"^[-*]\s", text, re.MULTILINE):
            for line in text.split("\n"):
                line = line.strip()
                if line.startswith(("- ", "* ")):
                    doc.add_paragraph(line[2:], style="List Bullet")
        elif re.match(r"^\d+\.\s", text, re.MULTILINE):
            for line in text.split("\n"):
                line = line.strip()
                match = re.match(r"^\d+\.\s+(.*)", line)
                if match:
                    doc.add_paragraph(match.group(1), style="List Number")
        else:
            doc.add_paragraph(text)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def export_to_pdf(content, title="Generated Document"):
    buffer = BytesIO()
    pdf = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        topMargin=inch,
        bottomMargin=inch,
        leftMargin=inch,
        rightMargin=inch,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Heading1"],
        fontSize=18,
        spaceAfter=16,
    )
    story = [Paragraph(_escape_pdf(title), title_style), Spacer(1, 12)]

    for block in _parse_blocks(content):
        level, text = _block_level(block)
        escaped = _escape_pdf(text)

        if level == 1:
            story.append(Paragraph(escaped, styles["Heading1"]))
        elif level == 2:
            story.append(Paragraph(escaped, styles["Heading2"]))
        elif level == 3:
            story.append(Paragraph(escaped, styles["Heading3"]))
        else:
            story.append(
                Paragraph(escaped.replace("\n", "<br/>"), styles["Normal"])
            )

        story.append(Spacer(1, 8))

    pdf.build(story)
    buffer.seek(0)
    return buffer
